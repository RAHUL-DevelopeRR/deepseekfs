"""Bounded, structure-aware extraction for the retrieval index.

Offsets refer to extracted section text (not PDF bytes). Pages are one-based.
Parsing is streamed so large documents do not become one giant string.
"""

from __future__ import annotations

import csv
import json
import os
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterator

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
VIDEO_EXTENSIONS = {".mp4", ".mkv", ".avi", ".mov", ".wmv", ".flv", ".webm"}
CHUNK_CHARS = 900
OVERLAP = 120
MAX_CHUNKS = 4096
MAX_CONTENT_BYTES = 64 * 1024 * 1024
PARSER_VERSION = "chunks-v2-token-bounded"


@dataclass
class Chunk:
    text: str
    section: str = ""
    page: int | None = None
    offset_start: int = 0
    offset_end: int = 0
    evidence_kind: str = "content"
    timestamp: str | None = None

    def to_dict(self) -> dict:
        return asdict(self)


def split_section(text: str, **metadata) -> Iterator[Chunk]:
    start = 0
    while start < len(text):
        end = min(start + CHUNK_CHARS, len(text))
        if end < len(text):
            # Prefer paragraph/line/word boundaries without dropping any text.
            for sep in ("\n\n", "\n", " "):
                boundary = text.rfind(sep, start + CHUNK_CHARS // 2, end)
                if boundary > start:
                    end = boundary + len(sep)
                    break
        if text[start:end].strip():
            yield Chunk(
                text=text[start:end], offset_start=start, offset_end=end, **metadata
            )
        if end == len(text):
            break
        start = max(start + 1, end - OVERLAP)


def _text_sections(path: Path) -> Iterator[tuple[str, dict]]:
    heading = "Text"
    buffer = ""
    block = 1
    with path.open(encoding="utf-8", errors="replace") as stream:
        while line := stream.readline(16000):
            if path.suffix.lower() == ".md" and re.match(r"^#{1,6}\s", line):
                if buffer:
                    yield buffer, {"section": f"{heading} / block {block}"}
                heading, buffer, block = line.strip(), "", 1
            buffer += line
            while len(buffer) >= 16000:
                yield buffer[:16000], {"section": f"{heading} / block {block}"}
                buffer = buffer[16000 - OVERLAP :]
                block += 1
        if buffer:
            yield buffer, {"section": f"{heading} / block {block}"}


def sections(path: Path) -> Iterator[tuple[str, dict]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        import fitz

        with fitz.open(path) as document:
            for number, page in enumerate(document, 1):
                yield (
                    page.get_text(sort=True),
                    {"page": number, "section": f"Page {number}"},
                )
    elif ext == ".docx":
        from docx import Document

        document = Document(path)
        heading = "Document"
        for number, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                heading = paragraph.text
            yield paragraph.text, {"section": f"{heading} / paragraph {number}"}
        for number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                yield (
                    " | ".join(cell.text for cell in row.cells),
                    {"section": f"Table {number}, row {row_number}"},
                )
    elif ext == ".pptx":
        from pptx import Presentation

        for number, slide in enumerate(Presentation(path).slides, 1):
            text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text.append(shape.text)
                if shape.has_table:
                    text.extend(
                        " | ".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )
            yield "\n".join(text), {"page": number, "section": f"Slide {number}"}
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                header = ""
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    text = " | ".join(
                        "" if value is None else str(value) for value in row
                    )
                    if number == 1:
                        header = text[:400]
                    if text.strip(" |"):
                        yield (
                            (f"Columns: {header}\n{text}" if number > 1 else text),
                            {"section": f"{sheet.title}, row {number}"},
                        )
        finally:
            book.close()
    elif ext == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as stream:
            header = ""
            for number, row in enumerate(csv.reader(stream), 1):
                text = " | ".join(row)
                if number == 1:
                    header = text[:400]
                yield (
                    (f"Columns: {header}\n{text}" if number > 1 else text),
                    {"section": f"Row {number}"},
                )
    elif ext == ".ipynb":
        with path.open(encoding="utf-8") as stream:
            notebook = json.load(stream)
        for number, cell in enumerate(notebook.get("cells", []), 1):
            if cell.get("cell_type") in {"code", "markdown"}:
                yield (
                    "".join(cell.get("source", [])),
                    {"section": f"Cell {number} ({cell['cell_type']})"},
                )
    elif ext in IMAGE_EXTENSIONS:
        from PIL import Image

        with Image.open(path) as image:
            yield (
                f"Image: {path.name}; format={image.format}; width={image.width}; height={image.height}",
                {"section": "Image metadata", "evidence_kind": "metadata"},
            )
            if os.getenv("NEURON_INDEX_IMAGE_OCR", "0") == "1":
                try:
                    import pytesseract

                    text = pytesseract.image_to_string(image, timeout=15)
                    yield text, {"section": "Image OCR", "evidence_kind": "ocr"}
                except Exception as exc:
                    from app.logger import logger

                    logger.warning("Image OCR unavailable for %s: %s", path, exc)
                    yield (
                        "OCR unavailable; image contents have not been read.",
                        {"section": "OCR status", "evidence_kind": "metadata"},
                    )
    elif ext in VIDEO_EXTENSIONS:
        yield (
            f"Video: {path.name}. Visual frames and audio have not been analyzed.",
            {"section": "Video metadata", "evidence_kind": "metadata"},
        )
        # Explicit sidecar evidence; no inference about unseen video frames.
        for suffix in (".srt", ".vtt"):
            sidecar = path.with_suffix(suffix)
            if not sidecar.is_file():
                continue
            if sidecar.stat().st_size > MAX_CONTENT_BYTES:
                yield (
                    f"Subtitle sidecar {sidecar.name} exceeds extraction limit.",
                    {"section": "Subtitle status", "evidence_kind": "metadata"},
                )
                continue
            timestamp, lines = None, []
            with sidecar.open(encoding="utf-8-sig", errors="replace") as stream:
                for line in stream:
                    if "-->" in line:
                        if lines and timestamp:
                            yield (
                                " ".join(lines),
                                {
                                    "section": sidecar.name,
                                    "timestamp": timestamp,
                                    "evidence_kind": "subtitle",
                                },
                            )
                        timestamp, lines = line.strip(), []
                    elif not line.strip():
                        if lines and timestamp:
                            yield (
                                " ".join(lines),
                                {
                                    "section": sidecar.name,
                                    "timestamp": timestamp,
                                    "evidence_kind": "subtitle",
                                },
                            )
                        timestamp, lines = None, []
                    elif line.strip() and timestamp:
                        lines.append(line.strip())
                if lines and timestamp:
                    yield (
                        " ".join(lines),
                        {
                            "section": sidecar.name,
                            "timestamp": timestamp,
                            "evidence_kind": "subtitle",
                        },
                    )
    elif ext in {".exe", ".msi", ".dll", ".lnk", ".doc", ".xls", ".ppt"}:
        yield (
            f"File: {path.name} (metadata only; content not extracted)",
            {"section": "File metadata", "evidence_kind": "metadata"},
        )
    elif ext in {".html", ".htm", ".xml"}:
        from html.parser import HTMLParser

        class Extractor(HTMLParser):
            def __init__(self):
                super().__init__()
                self.parts, self.hidden = [], 0

            def handle_starttag(self, tag, attrs):
                if tag in {"script", "style", "noscript"}:
                    self.hidden += 1

            def handle_endtag(self, tag):
                if tag in {"script", "style", "noscript"}:
                    self.hidden = max(0, self.hidden - 1)

            def handle_data(self, data):
                if not self.hidden:
                    self.parts.append(data)

        parser = Extractor()
        with path.open(encoding="utf-8", errors="replace") as stream:
            number = 0
            while data := stream.read(16000):
                parser.feed(data)
                number += 1
                yield " ".join(parser.parts), {"section": f"Markup block {number}"}
                parser.parts.clear()
    else:
        yield from _text_sections(path)


def parse_chunks(file_path: str) -> tuple[list[dict], bool]:
    path = Path(file_path)
    if (
        path.suffix.lower() not in VIDEO_EXTENSIONS
        and path.stat().st_size > MAX_CONTENT_BYTES
    ):
        return [
            Chunk(
                text=f"File: {path.name}. Content exceeds the 64 MiB extraction limit; metadata only.",
                section="Extraction limit",
                evidence_kind="metadata",
            ).to_dict()
        ], True
    chunks = []
    source = sections(path)
    try:
        for text, metadata in source:
            for chunk in split_section(text, **metadata):
                if len(chunks) >= MAX_CHUNKS:
                    return chunks, True
                chunks.append(chunk.to_dict())
    finally:
        source.close()
    return chunks, False
